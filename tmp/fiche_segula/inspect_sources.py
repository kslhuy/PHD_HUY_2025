from pathlib import Path
from pypdf import PdfReader
from docx import Document
from zipfile import ZipFile
import json, importlib.util

base=Path('presentation/rapport')
work=Path('tmp/fiche_segula')
r=PdfReader(base/'Final_PHD_these_2026.pdf')
pages=[p.extract_text() for p in r.pages]
(work/'these.txt').write_text('\n'.join(f'\n=== PAGE PDF {i+1} ===\n{t}' for i,t in enumerate(pages)),encoding='utf-8')
(work/'these_pages.json').write_text(json.dumps(pages,ensure_ascii=False),encoding='utf-8')
d=Document(base/'Fiche_Synthese_SEGULA_Capteur_Embarque_VAC_2.docx')
(work/'paragraphs.txt').write_text('\n'.join(f'{i} [{p.style.name}] {p.text}' for i,p in enumerate(d.paragraphs)),encoding='utf-8')
with ZipFile(base/'Fiche_Synthese_SEGULA_Capteur_Embarque_VAC_2.docx') as z:
    for n in z.namelist():
        if n.startswith('word/media/'):
            (work/Path(n).name).write_bytes(z.read(n))
print('PDF pages',len(pages),'DOCX paragraphs',len(d.paragraphs),'tables',len(d.tables),'images',len(d.inline_shapes))
print('fitz',importlib.util.find_spec('fitz'))
print('pdf2image',importlib.util.find_spec('pdf2image'))
for i,s in enumerate(d.sections):
    print('Section',i,'size',s.page_width.cm,s.page_height.cm,'margins',s.top_margin.cm,s.bottom_margin.cm,s.left_margin.cm,s.right_margin.cm)
