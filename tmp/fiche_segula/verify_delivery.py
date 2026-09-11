from pathlib import Path
from pypdf import PdfReader
from docx import Document
from hashlib import sha256
import json
w=Path('tmp/fiche_segula')
r=PdfReader(w/'final_render_v3/fiche.pdf')
old=PdfReader(w/'final_render_v2/fiche.pdf')
assert len(r.pages)==23
changes=[i+1 for i,(a,b) in enumerate(zip(r.pages,old.pages)) if a.extract_text()!=b.extract_text()]
print('Pages with changed text:',changes)
p=Path('presentation/rapport/Fiche_Synthese_SEGULA_Capteur_Embarque_VAC_3_enrichie.docx')
d=Document(p)
text='\n'.join(x.text for x in d.paragraphs)
assert all(s in text for s in ['Robust KalmanNet','retour arrière','100 Hz','50 Kio','168 octets','Exemple chiffré de dimensionnement uniquement.'])
source=Path('presentation/rapport/Fiche_Synthese_SEGULA_Capteur_Embarque_VAC_2.docx')
assert sha256(source.read_bytes()).hexdigest()=='744eb795b22cc76b9dfeee0a9e52f6f80dd808bb86b6408ae937531a5da82306'
assert len(d.inline_shapes)==15
print('Source unchanged; final valid; images 15; tables',len(d.tables),'pages',len(r.pages))
report={'pages':23,'source_unchanged':True,'final_sha256':sha256(p.read_bytes()).hexdigest(),'images':15,'tables':len(d.tables),'changed_pages_last_pass':changes,'render':'Microsoft Word PDF export with local printer; source DOCX structural reconstruction required','qa':'All 23 pages visually inspected; final text corrections on pages 22 and 23'}
(w/'verification.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
