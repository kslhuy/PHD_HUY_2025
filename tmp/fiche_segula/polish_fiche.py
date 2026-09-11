from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
p=Path('presentation/rapport/Fiche_Synthese_SEGULA_Capteur_Embarque_VAC_3_enrichie.docx')
d=Document(p)
for st in d.styles:
    for border in list(st.element.xpath('.//w:pBdr')):border.getparent().remove(border)
for par in d.paragraphs:
    for border in list(par._p.xpath('./w:pPr/w:pBdr')):border.getparent().remove(border)
    if par.text.startswith('L’opération établit une chaîne'):
        par.text='Le démonstrateur établit une chaîne embarquée d’acquisition, d’estimation et de communication. Le bilan distingue les fonctions opérationnelles des extensions à valider sur carte et des mesures de ressources à compléter.'
    elif par.text.startswith('La contribution principale est la réalisation'):
        par.text='La contribution principale est la réalisation du capteur : pilotes, EKF périodique, supervision et intégration ROS 2. Les essais à venir doivent relier les gains sous attaque au coût sur la carte, puis vérifier leur reproductibilité sur la carte dédiée.'
    elif par.text.startswith('Ce projet s’inscrit pour partie'):
        par.text='Le projet associe SEGULA Technologies et le CRAN dans le cadre de la thèse CIFRE de NGUYEN Quang Huy. Le CRAN apporte son expertise en électronique, estimation et commande, ainsi que sa plateforme de validation.'
        par.paragraph_format.keep_together=True
for t in d.tables:
    for c in t.rows[0].cells:
        for pp in c.paragraphs:pp.paragraph_format.keep_with_next=True
ann=d.tables[-1]
for c,text in zip(ann.rows[0].cells,['Classe R&D','Domaines et mots clés','Présentation synthétique','Indicateurs de recherche','Observations','Thèse CIFRE']):c.text=text
widths=[2.0,4.0,8.6,6.35,2.0,2.0]
for col,w in zip(ann.columns,widths):col.width=Cm(w)
for ri,row in enumerate(ann.rows):
    for c,w in zip(row.cells,widths):
        c.width=Cm(w)
        for pp in c.paragraphs:
            pp.paragraph_format.space_before=Pt(3);pp.paragraph_format.space_after=Pt(3)
            for r in pp.runs:r.font.size=Pt(9);r.bold=(ri==0)
d.save(p)
print('Polished',p)
