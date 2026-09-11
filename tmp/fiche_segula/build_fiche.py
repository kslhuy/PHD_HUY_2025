from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import hashlib, json, re

ROOT=Path.cwd()
SRC=ROOT/'presentation/rapport/Fiche_Synthese_SEGULA_Capteur_Embarque_VAC_2.docx'
OUT=ROOT/'presentation/rapport/Fiche_Synthese_SEGULA_Capteur_Embarque_VAC_3_enrichie.docx'
doc=Document(SRC)
old=list(doc.paragraphs)

def replace(i,text):
    p=old[i]; p.clear(); p.add_run(text); return p

def add_before(i,text,style='Normal',boldlead=None):
    p=old[i].insert_paragraph_before(style=style)
    if boldlead and text.startswith(boldlead):
        p.add_run(boldlead).bold=True; p.add_run(text[len(boldlead):])
    else:p.add_run(text)
    return p

def heading(i,text,level=3):return add_before(i,text,f'Heading {level}')

def table_before(i,headers,rows,widths):
    t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,text,w in zip(t.rows[0].cells,headers,widths):c.text=text;c.width=Cm(w)
    for row in rows:
        for c,text,w in zip(t.add_row().cells,row,widths):c.text=str(text);c.width=Cm(w)
    for col,w in zip(t.columns,widths):col.width=Cm(w)
    old[i]._p.addprevious(t._tbl)
    format_table(t)
    return t

def format_table(t):
    pr=t._tbl.tblPr
    borders=pr.find(qn('w:tblBorders'))
    if borders is None:borders=OxmlElement('w:tblBorders');pr.append(borders)
    for side in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    for ri,row in enumerate(t.rows):
        rp=row._tr.get_or_add_trPr()
        for e in list(rp):
            if e.tag==qn('w:trHeight'):rp.remove(e)
        cant=OxmlElement('w:cantSplit');rp.append(cant)
        if ri==0:
            hdr=OxmlElement('w:tblHeader');rp.append(hdr)
        for c in row.cells:
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr()
            mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'95');el.set(qn('w:type'),'dxa');mar.append(el)
            cp.append(mar)
            for e in list(cp):
                if e.tag==qn('w:shd'):cp.remove(e)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCEAF2' if ri==0 else ('F5F8FA' if ri%2==0 else 'FFFFFF'));cp.append(sh)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                p.paragraph_format.line_spacing=1.05
                for r in p.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor(0,0,0);r.bold=(ri==0)
    return t

# Correct the scope of claims while keeping the corporate structure and figures.
replace(24,'Capteur embarqué intelligent pour véhicules autonomes connectés')
old[24].style=doc.styles['Title']
add_before(25,'Implantation sur carte et extension vers une estimation résiliente','Subtitle')
replace(34,"La présente opération porte sur le passage des méthodes d’estimation à une exécution embarquée à 100 Hz. Le démonstrateur associe une carte de développement, des capteurs, un firmware temps réel et une passerelle vers ROS 2 ; il prépare le transfert vers la carte électronique dédiée du projet. L’EKF constitue la brique locale déjà décrite comme implantée. La confiance distribuée, le retour arrière et Robust KalmanNet prolongent cette architecture et font l’objet ci-après d’une proposition de portage et de validation sur carte.")
replace(45,"La problématique centrale est l’embarquabilité : traduire un algorithme d’estimation en un calcul discret compatible avec les ressources mémoire, les bus et l’ordonnancement du microcontrôleur. Pour une période de 10 ms, il faut maîtriser l’âge des mesures, le temps de calcul, les interruptions et les échanges entre cœurs. La cadence nominale doit être complétée par une mesure de la gigue et des dépassements d’échéance avant de conclure au respect du temps réel dur.")
replace(57,"Les travaux antérieurs cités dans l’opération concernent les observateurs non linéaires et à entrées inconnues, puis l’estimation distribuée et la confiance pour la conduite en convoi. Le support de thèse [14] développe plus particulièrement les observateurs local et distribué, la confiance quantitative, le retour arrière après détection tardive et Robust KalmanNet. L’articulation avec la carte repose sur une estimation locale exploitable, des données datées et des indicateurs de qualité accessibles au niveau du firmware.")
replace(68,'[3] NGUYEN Quang Huy, travaux de thèse CIFRE VEHALSECU, CRAN et SEGULA Technologies, 2023–2026. Intitulé de la thèse indiqué en section 2 ; support algorithmique détaillé en [14].')
add_before(79,'[14] NGUYEN Quang Huy, Final_PHD_these_2026.pdf, support de présentation fourni, 46 pages. Observateurs : pp. PDF 4–7 ; confiance : pp. 8–25 ; simulations : pp. 26–28 ; retour arrière : pp. 29–33 ; Robust KalmanNet : pp. 34–46. La numérotation affichée dans le support va de 30 à 75. Les renvois ci-après utilisent les pages du PDF.')
replace(80,"L’intérêt de l’opération réside dans l’intégration d’une chaîne d’estimation ouverte sur une cible contrainte, avec accès aux mesures, innovations et états de fonctionnement. La qualité de l’estimation, la robustesse aux données incohérentes et le coût embarqué doivent être évalués ensemble. Les résultats algorithmiques de [14] motivent les extensions de résilience, mais leur gain sur la carte doit être établi par des essais comparatifs à ressources et scénarios identiques.")
replace(91,"L’opération a établi un démonstrateur d’estimation embarquée intégré au véhicule d’expérimentation. L’apport matériel et logiciel porte sur l’acquisition, le calcul de l’EKF, la supervision et l’export des estimations. Les contributions rapportées sont les suivantes :")
replace(94,"Caractérisation rapportée sur une séquence dynamique de 20 s : RMSE de vitesse de 0,011 m/s, de cap de 0,43° et de position horizontale de 0,39 m. Ces valeurs décrivent cette séquence et ne constituent pas une précision garantie en tout environnement.")
replace(95,"Caractérisation du résidu normalisé sur la séquence étudiée : maximum rapporté de 2,12 et aucune fausse alarme observée. Ce résultat fournit un repère pour le réglage du diagnostic, à compléter par une campagne plus longue et des injections d’attaque.")
replace(96,"Mise en œuvre de mécanismes de robustesse : calibration du magnétomètre, garde sur l’innovation magnétique, correction à vitesse nulle et estimation du biais gyroscopique. Leur contribution doit être isolée par des comparaisons avec et sans chaque mécanisme.")
replace(99,"Définition précise du résidu et calibration du seuil d’alarme sur un ensemble sain, puis validation sur des données indépendantes et sous attaque.")
replace(102,"Portage vers la carte dédiée avec vérification des broches, horloges, interruptions, mémoires partagées, alimentation et comportement temporel ; évaluation progressive de la confiance distribuée, du retour arrière et de Robust KalmanNet.")
replace(104,"La modularité des pilotes et de l’estimateur facilite le transfert vers la carte dédiée. La réutilisation de composants capteurs communs réduit les adaptations, mais le portage comprend aussi les bus, les horloges, les interruptions, la synchronisation et les essais de non-régression. La reproductibilité repose sur l’archivage conjoint de la version du firmware, du câblage, des paramètres de calibration, des matrices du filtre et des journaux horodatés. Chaque indicateur doit rester associé au scénario, à la référence utilisée et à la durée de l’essai.")
replace(109,"Les premiers travaux ont porté sur la carte électronique dédiée CRAN / SEGULA : deux microcontrôleurs assurent les fonctions de navigation et de communication, avec capteurs inertiel et magnétique, récepteur GNSS, mémoires et interfaces de liaison. L’étude des schémas et des connecteurs a permis d’identifier les bus et les ressources disponibles. Cette carte dédiée doit être distinguée de la carte d’évaluation à microcontrôleur double cœur utilisée pour poursuivre le développement.")
replace(116,"La plateforme de substitution est une carte de développement STM32H745-DISCO, identifiée dans le schéma de la figure 3, associée à des modules capteurs. Elle permet de développer les pilotes, le traitement et les échanges entre cœurs malgré les défaillances de la carte dédiée. La logique applicative est réutilisable ; son transfert nécessite une nouvelle vérification des interfaces et du comportement sur la carte cible.")
replace(125,"L’architecture distingue la carte dédiée à deux microcontrôleurs, la plateforme de développement STM32 à deux cœurs et la passerelle ESP32. Sur le démonstrateur, le cœur secondaire assure l’acquisition ; le cœur principal exécute l’estimation, l’affichage et la préparation de la télémétrie. L’ESP32 relaie les données sur le réseau. Cette séparation réduit le couplage avec la communication, à condition que les files, interruptions et accès partagés restent bornés et que l’estimateur n’attende jamais une transmission réseau.")
replace(131,"La temporisation à référence absolue évite l’accumulation de dérive liée à une attente relative. Elle ne suffit pas à garantir le respect des échéances lorsque la charge augmente. Les pilotes décrits assurent la lecture SPI de l’inertiel, la réception GNSS par interruption et la lecture I²C du magnétomètre. La validation temporelle doit mesurer le début et la fin des tâches, les blocages sur ressources partagées et les éventuels retards de réveil.")
replace(157,'Figure 8 : suivi de vitesse et erreur sur la séquence de caractérisation de 20 s ; le RMSE est un indicateur moyen, pas une borne de l’erreur instantanée')
replace(155,"Caractérisation quantitative de l’estimateur. La fiche rapporte une séquence de 20 s à 100 Hz, avec accélérations, décélérations, freinage avec glissement, biais gyroscopique injecté de 12 mrad/s et perturbations magnétiques. Les valeurs des tableaux 3 et 4 sont conservées pour cette séquence. Leur utilisation comme mesures de performance matérielle exige de rattacher les journaux à la carte exécutant le calcul et de préciser la référence de vitesse, de cap et de position.")
replace(160,"Sur cette séquence, le RMSE de vitesse représente 11 mm/s. Le maximum rapporté, 0,063 m/s, montre que l’erreur instantanée peut dépasser le RMSE. Le R² élevé indique une bonne concordance avec la référence utilisée, sans renseigner à lui seul la robustesse sous attaque. La durée de 20 s et la référence d’essai doivent accompagner ces chiffres lors de leur réutilisation.")
replace(162,"En fonctionnement nominal, le maximum rapporté de 2,12 décrit l’amplitude observée du résidu normalisé sur la séquence ; il ne constitue ni un plancher de bruit ni une borne universelle. Un dépassement de seuil signale une incohérence qui peut provenir d’une attaque, d’un défaut, d’un mauvais modèle, d’un retard ou d’une perturbation de l’environnement. La décision doit croiser les voies de mesure, la persistance de l’écart et la disponibilité des capteurs. L’absence de fausse alarme sur 20 s ne suffit pas à estimer un taux de fausses alertes en exploitation.")
replace(167,"L’opération établit une chaîne d’acquisition, d’estimation et de communication sur une plateforme embarquée de développement. Elle constitue une étape vers la carte dédiée et vers la résilience distribuée. Le bilan distingue les fonctions décrites comme opérationnelles des extensions algorithmiques et des mesures de ressources restant à valider sur cible.")
replace(169,"La contribution principale est la réalisation du capteur embarqué et de sa chaîne de validation : pilotes, exécution périodique de l’EKF, mécanismes de robustesse, supervision et intégration ROS 2. Les résultats rapportés donnent une référence nominale. L’étape suivante consiste à relier chaque gain sous attaque à son coût sur la carte, puis à reproduire les essais sur la carte dédiée.")

# General algorithm overview from the actual PDF, without claiming MCU deployment.
heading(55,'Algorithmes de la thèse et rôle dans le capteur',3)
add_before(55,"Le capteur local fournit l’état estimé du véhicule et la qualité des mesures. L’observateur distribué exploite ensuite les échanges entre véhicules pour construire une estimation du convoi. Les méthodes de confiance agissent sur les données utilisées par cet observateur ; le retour arrière corrige une contamination déjà propagée ; Robust KalmanNet traite la fiabilité des mesures au niveau de l’estimation locale [14].")
heading(55,'Observateurs local et distribué',4)
add_before(55,"Le modèle du support décrit un véhicule de type bicyclette par sa position plane, sa vitesse, son accélération et son cap. L’observateur local corrige la prédiction à partir des mesures disponibles. L’observateur distribué combine une estimation propre, des estimations reçues et une contribution locale, avec des poids de consensus non négatifs. Ce modèle à cinq états se distingue de l’EKF de navigation à six états de la carte, qui comprend les vitesses Nord et Est ainsi qu’un biais gyroscopique. Leur couplage nécessite une conversion explicite des états, repères, unités et dates [14, pp. 4–7].")
heading(55,'Confiance quantitative et sélection des informations',4)
add_before(55,"Un indicateur de validité filtre d’abord les messages selon leur authentification et leur fraîcheur. La confiance locale LT vérifie la cohérence physique de la vitesse, de la distance, de l’accélération et du cap. La confiance distribuée DT confronte les estimations reçues à celle du véhicule hôte et aux mesures relatives locales, puis vérifie la cohérence interne de l’estimation. Les écarts sont normalisés ; plusieurs facteurs utilisent une décroissance exponentielle de type exp(−D) [14, pp. 8–18].")
add_before(55,"Le score instantané est le produit LT × DT. Un historique à cinq niveaux lisse la confiance afin de limiter les changements dus au bruit ou à une perte brève de paquets. Le vecteur de confiance étend l’opinion aux autres véhicules, directement ou par propagation à un saut. Un seuil sélectionne les contributions admises dans le consensus. L’authentification d’un émetteur ne prouve pas la cohérence physique de son contenu ; ces deux vérifications se complètent [14, pp. 20–25].")
heading(55,'Retour arrière après détection tardive',4)
add_before(55,"Une donnée falsifiée peut contaminer l’estimation d’un voisin honnête avant d’être détectée. Le retour arrière conserve un historique des états, entrées et contributions ; après détection, il repart d’un état antérieur et rejoue les mises à jour en excluant les sources reconnues comme compromises. L’intérêt sur la carte est de réparer l’état courant avant de poursuivre la prédiction. L’efficacité dépend de la profondeur d’historique, de la date de contamination et du délai disponible pour le recalcul [14, pp. 29–33].")
heading(55,'Robust KalmanNet pour les mesures corrompues',4)
add_before(55,"Robust KalmanNet conserve une prédiction issue du modèle physique et apprend la correction des mesures. Ses entrées comprennent l’innovation, les variations de prédiction et de mesure ainsi que les indicateurs de disponibilité. Un réseau produit un masque de fiabilité par voie ; une structure récurrente GRU contribue au gain de correction appris. Le masque atténue les innovations des voies jugées corrompues, notamment la position GNSS et la vitesse. L’apprentissage et la validation sont réalisés hors carte ; le portage envisagé concerne l’inférence avec des poids figés [14, pp. 34–44].")
add_before(55,"Les courbes de [14, pp. 45–46] montrent, sur l’essai présenté, une réduction des écarts à la référence « Clean Ref EKF » lors de perturbations et une adaptation des masques et gains. Cette référence est un estimateur de comparaison, pas une vérité terrain indépendante. Ces figures n’établissent pas le temps d’inférence, la mémoire consommée ou un gain chiffré de Robust KalmanNet sur la carte STM32.")
table_before(55,['Brique','Lien avec la carte','Niveau présenté'],[
 ['EKF de navigation','Acquisition et fusion locale à 100 Hz','Implantation décrite dans la fiche'],
 ['Confiance et consensus','Filtrer et pondérer les contributions V2V','Algorithmes et essais dans [14] ; portage proposé'],
 ['Retour arrière','Historique et recalcul après alerte','Mécanisme dans [14] ; coût cible à mesurer'],
 ['Robust KalmanNet','Masque et gain appris pour la correction locale','Résultats comparatifs dans [14] ; inférence cible à qualifier']
],[3.1,6.0,6.8])

# Board-specific additions, placed alongside the existing implementation narrative.
heading(132,'Synchronisation et qualité des données sur la carte',4)
add_before(132,"Spécification de consolidation du firmware. Chaque pilote doit produire une mesure associée à un horodatage d’acquisition, un compteur et un état de validité. À chaque cycle de 10 ms, l’EKF récupère un instantané cohérent des dernières mesures. Une correction GNSS ou magnétique ne doit être appliquée qu’une fois par nouvel échantillon ; répéter une même mesure à 100 Hz conduirait à surestimer l’information disponible.",boldlead='Spécification de consolidation du firmware.')
add_before(132,"La mémoire partagée doit contenir des structures de taille fixe. Un double tampon avec numéro de version, ou une courte section critique, permet d’éviter la lecture d’un mélange de deux acquisitions. Le protocole retenu devra être validé avec les mécanismes de cohérence mémoire propres à la cible. Les temps d’attente sur bus doivent être bornés ; la télémétrie et l’affichage récupèrent une copie de l’état sans bloquer l’estimateur.")
add_before(132,"Les contrôles d’entrée portent sur les unités, le repère Nord/Est, la convention de cap, la validité GNSS et l’âge de la donnée. En cas d’absence de nouvelle mesure, la prédiction continue selon le mode dégradé prévu et l’incertitude doit être propagée. Un journal distingue mesure manquante, mesure rejetée et perte de trame réseau, afin d’attribuer correctement les dégradations de performance.")
heading(137,'Du calcul matriciel au code de la carte',4)
add_before(137,"Pour l’EKF à six états, la covariance est une matrice 6 × 6. La prédiction propage l’état et cette covariance. Une correction scalaire calcule l’innovation r = z − h(x prédit), sa variance S, puis le gain à partir de la covariance et de S. Cette organisation évite l’inversion d’une grande matrice de mesures, sous l’hypothèse de bruits de mesure indépendants ou préalablement décorrélés.")
add_before(137,"Pour le portage, les tailles de matrices doivent être fixées à la compilation, les buffers préalloués et les paramètres de bruit versionnés. Les gardes numériques doivent contrôler S strictement positif, les valeurs non finies, la symétrie de la covariance et le repli des écarts de cap dans l’intervalle angulaire choisi. Le calcul en précision simple devra être comparé à une référence hors carte avant toute optimisation.")
add_before(137,"Le diagnostic doit distinguer le résidu standardisé r/√S du carré d’innovation normalisé r²/S, ou NIS en dimension scalaire. Un test du khi-deux s’applique au NIS sous ses hypothèses statistiques ; il ne faut pas comparer directement ces différentes grandeurs au même seuil. Une persistance sur plusieurs mises à jour et un seuil calibré sur des données indépendantes complètent la décision.")

heading(148,'Portage des fonctions de résilience sur la carte',3)
add_before(148,"Architecture proposée. L’EKF conserve la priorité à 100 Hz. Une tâche de supervision traite les informations disponibles à cadence plus faible, puis publie des décisions de validité et de pondération. Le déploiement doit commencer par un mode d’observation où les nouveaux algorithmes calculent leurs décisions sans modifier l’état utilisé par le véhicule ; le couplage est activé après comparaison des journaux.",boldlead='Architecture proposée.')
heading(148,'Confiance et observateur distribué',4)
add_before(148,"Une implantation initiale à 10 Hz, alignée sur la télémétrie, est proposée pour la confiance. Pour un convoi de quatre véhicules, chaque hôte gère au plus trois voisins directs : états datés, compteurs, scores LT et DT, historique et poids. Le module de validité doit vérifier la fraîcheur et l’authenticité effective des messages. Le compteur de la liaison UDP actuelle aide à détecter les pertes ; il ne constitue pas une authentification.")
add_before(148,"Les entrées utiles comprennent la position, la vitesse, le cap, l’incertitude et, lorsque disponibles, des mesures relatives de distance ou de vitesse. Sans ces mesures relatives, certains tests DT du support ne sont pas réalisables tels quels. Le firmware devra borner le nombre de voisins, gérer les données absentes et conserver des poids non négatifs dont la somme est cohérente avec le terme propre de l’observateur. Le passage de 100 à 10 Hz impose de recalculer les paramètres temporels du lissage et de vérifier le délai de détection.")
heading(148,'Historique et retour arrière borné',4)
add_before(148,"Le retour arrière nécessite un tampon circulaire contenant les états sauvegardés, les entrées, les durées d’échantillonnage, les contributions reçues et leurs poids. La profondeur doit couvrir le retard de détection visé. Le recalcul s’exécute sur une copie de l’état, avec un quota de travail par cycle ; l’état corrigé n’est publié qu’après rattrapage du temps courant. Si le début de contamination précède l’historique conservé, la réparation complète n’est pas assurée. Il faut alors définir une réinitialisation ou un repli vers l’estimation locale.")
heading(148,'Inférence de Robust KalmanNet',4)
add_before(148,"Le portage doit conserver les normalisations des entrées, les poids et l’état récurrent du réseau appris. L’entraînement, décrit en deux phases dans [14], reste hors carte. Un essai d’inférence mesure le coût des couches, des buffers et des activations avant de viser une cadence de 100 Hz. La quantification ou la réduction du réseau constituent des pistes d’optimisation, à accepter seulement si les erreurs et la détection restent conformes sur un jeu d’essai indépendant.")
add_before(148,"La stratégie de supervision doit prévoir un retour à l’EKF de référence si l’inférence dépasse son délai, produit une valeur non finie ou perd sa cohérence. Le modèle bicyclette de Robust KalmanNet utilise des commandes et paramètres qui doivent être accessibles et identifiés sur le véhicule. Le remplacement direct de l’EKF à six états n’est donc pas automatique : l’adaptateur d’états et le maintien de l’estimation du biais doivent être conçus et testés.")

heading(148,'Budget de ressources et indicateurs de la carte',3)
add_before(148,"Exemple chiffré de dimensionnement uniquement. Les allocations ci-dessous sont des hypothèses de travail pour préparer une campagne de mesure ; elles ne sont pas des temps observés sur STM32. Pour un cycle de 10 ms, l’enveloppe critique proposée est de 2,0 ms, soit 20 % du cycle. Le solde de 8,0 ms est une réserve théorique avant prise en compte de toutes les interruptions, attentes et contentions.",boldlead='Exemple chiffré de dimensionnement uniquement.')
table_before(148,['Traitement dans un cycle','Allocation illustrative','Vérification sur cible'],[
 ['Copie et contrôle des mesures','0,5 ms','Âge des données et contention entre cœurs'],
 ['Prédiction et corrections EKF','0,9 ms','Cas avec toutes les corrections disponibles'],
 ['Supervision locale et résidus','0,4 ms','Charge maximale et traitement des rejets'],
 ['Préparation de la télémétrie','0,2 ms','Aucune attente de fin de transmission'],
 ['Total de l’enveloppe critique','2,0 ms sur 10 ms','Ne comprend pas le réseau appris ni le rejeu']
],[6.0,3.2,6.7])
add_before(148,"Exemple mémoire calculé. En précision simple, les 6 valeurs d’état et les 36 valeurs de covariance occupent 42 × 4 = 168 octets, hors matrices intermédiaires, piles et pilotes. Pour illustrer un historique, une structure supposée de 256 octets par instant et une profondeur de 200 instants nécessitent 51 200 octets, soit 50 Kio, et couvrent 2 s à 100 Hz. La structure réelle du retour arrière distribué peut être plus grande : son contenu et son nombre de voisins doivent être inclus dans le calcul.",boldlead='Exemple mémoire calculé.')
add_before(148,"Mesures à produire sur carte. Instrumenter les tâches par compteur de cycles ou signaux GPIO, puis relever les temps moyen, au 99e centile et maximum observé, la gigue, les dépassements de 10 ms, la marge de pile, la RAM, la Flash et les pertes de buffers. Le maximum observé sur une campagne est un résultat expérimental ; il ne constitue pas à lui seul une borne théorique de pire cas. Le courant moyen et les pointes de consommation peuvent compléter le bilan matériel.",boldlead='Mesures à produire sur carte.')

heading(161,'Lien entre choix embarqués et performances',4)
table_before(161,['Choix dans la carte','Indicateur associé','Lecture du résultat'],[
 ['Biais gyroscopique et correction magnétique','RMSE cap 0,43° ; écart au biais 0,055 mrad/s','Valeurs rapportées sur la séquence de 20 s'],
 ['Garde magnétique et correction à vitesse nulle','23 rejets ; 6 activations ZUPT','Vérifier les instants d’activation et les faux rejets'],
 ['Prédiction 100 Hz et mesures plus lentes','RMSE vitesse 0,011 m/s ; position 0,39 m','Isoler le gain par comparaison avec et sans fusion'],
 ['Passerelle et compteur de séquence','10 trames/s ; pertes de l’ordre de 2 %','Mesure de transport rapportée, distincte de la précision'],
 ['Ordonnancement et mémoire partagée','Gigue, latence, RAM et dépassements','Mesures cible à ajouter au bilan']
],[5.1,5.1,5.7])
add_before(161,"La campagne doit séparer les essais extérieurs, les tests du transport et la séquence de caractérisation de l’estimateur. Les 196 corrections GNSS sur 20 s représentent environ 9,8 corrections/s ; leur compatibilité avec une cadence de 10 Hz doit être contrôlée sur les horodatages. L’écart-type de cap inférieur à 0,2° décrit une stabilité après calibration, tandis que le RMSE de cap de 0,43° est un écart à une référence : ces indicateurs ne mesurent pas la même propriété.")

heading(166,'Campagne proposée pour relier résilience et coût embarqué',3)
add_before(166,"Le support [14, p. 26] fournit un scénario de simulation de quatre véhicules avec une attaque du véhicule 1 entre 10 et 15 s : biais de position de −5 m, biais de vitesse de −2 m/s, perturbations intermittentes et pertes avec une probabilité de 0,5. Ces paramètres servent de point de départ à un rejeu sur carte. Ils décrivent des entrées d’essai, pas des performances obtenues par le démonstrateur.")
table_before(166,['Essai proposé','Comparaison à réaliser','Mesures attendues'],[
 ['Nominal et charge maximale','EKF seul puis extensions activées','RMSE, gigue, RAM, échéances et fausses alertes'],
 ['Biais GNSS −5 m de 10 à 15 s','EKF et confiance locale ou RKNet','Pic d’erreur, délai de rejet, retour au nominal'],
 ['Biais de vitesse −2 m/s','Fusion nominale puis correction robuste','Erreur vitesse, distance estimée et faux rejets'],
 ['Pertes de messages à 50 %','Confiance avec gestion des absences','Continuité locale, âge des données, reprise'],
 ['Propagation et détection tardive','Avec et sans retour arrière','Erreur résiduelle, profondeur et temps de rejeu'],
 ['Saturation de communication','Télémétrie et affichage sollicités','Débordements, latence et tenue du cycle de 10 ms']
],[4.5,5.5,5.9])
add_before(166,"Pour chaque essai, conserver les mêmes données d’entrée et la même référence entre méthodes, puis répéter les scénarios stochastiques avec des graines consignées. Rapporter les erreurs hors attaque, pendant l’attaque et pendant la récupération, ainsi que le délai de détection, le taux de détection, les fausses alertes par durée d’observation et le coût processeur. Un rejeu sur carte établit le coût matériel du calcul ; un essai sur véhicule complète l’analyse des effets des capteurs et de l’environnement.")

# Fix the existing hardware and summary tables without inventing component specifications.
t=next(t for t in doc.tables if t.cell(0,0).text=='Fonction')
for row in t.rows[1:]:
    if row.cells[0].text.startswith('Navigation'):
        row.cells[1].text='STM32H745 du démonstrateur, cœur principal'
    elif row.cells[0].text=='Communication':
        row.cells[0].text='Acquisition';row.cells[1].text='STM32H745 du démonstrateur, cœur secondaire';row.cells[2].text='Mémoire partagée'
# Locate tables by their content because added tables change their indices.
for t in doc.tables:
    if t.cell(0,0).text=='Objectif de l’opération':
        for row in t.rows[1:]:
            if 'Observateur embarqué' in row.cells[0].text:row.cells[2].text='EKF décrit à 100 Hz sur démonstrateur ; profilage temporel à compléter'
            if 'Plan de tests' in row.cells[0].text:row.cells[1].text='Réalisé en partie';row.cells[2].text='Essais rapportés ; durée, référence et journaux à rattacher aux indicateurs'
        for vals in [('Carte dédiée','À réaliser','Remise en service, portage et nouvelle validation'),('Confiance et retour arrière','Proposé','Algorithmes [14] à profiler et éprouver sur la carte'),('Robust KalmanNet sur carte','À qualifier','Inférence, ressources et comparaison à l’EKF')]:
            cells=t.add_row().cells
            for c,v in zip(cells,vals):c.text=v

# Update the four-sentence executive description in the existing annex.
annex=doc.tables[-1]
if len(annex.columns)>=4:
    for row in annex.rows:
        for c in row.cells:
            if c.text.startswith("L'opération porte"):
                c.text=("L’opération développe un capteur embarqué intelligent associant des pilotes de capteurs, un EKF à six états exécuté à 100 Hz sur une plateforme STM32 et une liaison ESP32 vers ROS 2. "
                "La carte dédiée a été étudiée et ses défaillances ont conduit au développement sur une cible intermédiaire. "
                "La fiche rapporte sur une séquence de 20 s des RMSE de 0,011 m/s en vitesse, 0,43° en cap et 0,39 m en position, à rattacher à la référence et aux journaux d’essai. "
                "Les extensions issues du support de thèse concernent la confiance distribuée, le retour arrière et Robust KalmanNet ; leur validation sur carte doit quantifier le gain sous attaque ainsi que les temps de calcul et la mémoire.")

# Keep source page furniture and styles; make headings black and remove decorative title rules.
for name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3','Heading 4','Titre pdg','Titre pdg 1']:
    st=doc.styles[name];st.font.color.rgb=RGBColor(0,0,0)
    if st.element.rPr is not None:
        for color in st.element.rPr.findall(qn('w:color')):
            for key in list(color.attrib):
                if key!=qn('w:val'):del color.attrib[key]
    pp=st.element.pPr
    if pp is not None:
        for b in list(pp):
            if b.tag==qn('w:pBdr'):pp.remove(b)
doc.styles['Title'].font.name='Arial';doc.styles['Title'].font.size=Pt(19);doc.styles['Title'].font.bold=True
doc.styles['Subtitle'].font.name='Arial';doc.styles['Subtitle'].font.size=Pt(12)
for p in doc.paragraphs:
    if p.style.name.startswith('Heading') or p.style.name in ['Title','Subtitle','Titre pdg','Titre pdg 1']:
        for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
        p.paragraph_format.keep_with_next=True
    if p.style.name=='Caption':p.paragraph_format.keep_with_next=False;p.paragraph_format.keep_together=True
for t in doc.tables:
    format_table(t)

# Remove obsolete bookmarks that point to the old generated text and refresh the TOC in Word.
for el in list(doc.element.xpath('.//w:bookmarkStart | .//w:bookmarkEnd')):
    el.getparent().remove(el)
uf=doc.settings.element.find(qn('w:updateFields'))
if uf is None:uf=OxmlElement('w:updateFields');doc.settings.element.append(uf)
uf.set(qn('w:val'),'true')
doc.core_properties.title='Capteur embarqué intelligent pour véhicules autonomes connectés'
doc.core_properties.subject='Implantation sur carte, estimation et résilience'
doc.save(OUT)
print(OUT)
print('SHA256 source',hashlib.sha256(SRC.read_bytes()).hexdigest())
print('paragraphs',len(doc.paragraphs),'tables',len(doc.tables))
