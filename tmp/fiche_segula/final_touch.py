from docx import Document
from docx.shared import Pt
from pathlib import Path
p=Path('presentation/rapport/Fiche_Synthese_SEGULA_Capteur_Embarque_VAC_3_enrichie.docx')
d=Document(p)
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            if c.text=='Principe établi et plancher de bruit caractérisé ; validation sous attaque à mener':
                c.text='Résidu nominal caractérisé ; validation quantitative sous attaque à mener'
                for par in c.paragraphs:
                    par.paragraph_format.space_before=Pt(3);par.paragraph_format.space_after=Pt(3)
                    for r in par.runs:r.font.size=Pt(9)
c=d.tables[-1].rows[0].cells[4]
c.text='Notes'
for par in c.paragraphs:
    par.paragraph_format.space_before=Pt(3);par.paragraph_format.space_after=Pt(3);par.paragraph_format.keep_with_next=True
    for r in par.runs:r.font.size=Pt(9);r.bold=True
d.save(p)
