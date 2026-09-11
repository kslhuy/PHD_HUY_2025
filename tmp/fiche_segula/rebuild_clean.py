from pathlib import Path
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START,WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
from PIL import Image
from io import BytesIO
import re

root=Path.cwd();work=root/'tmp/fiche_segula'
out=root/'presentation/rapport/Fiche_Synthese_SEGULA_Capteur_Embarque_VAC_3_enrichie.docx'
src=Document(out);dst=Document()
sec=dst.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Cm(2.5)
sec.header_distance=Cm(1.0);sec.footer_distance=Cm(1.25)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Heading 4','Caption','List Paragraph']:
    s=dst.styles[name];s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_after=Pt(6)
    s.paragraph_format.line_spacing=1.08
    if s.element.rPr is not None:
        s.element.rPr.rFonts.set(qn('w:eastAsia'),'Arial')
dst.styles['Normal'].font.size=Pt(10.5)
dst.styles['Title'].font.size=Pt(20);dst.styles['Title'].font.bold=True
dst.styles['Subtitle'].font.size=Pt(13)
for n,size in [(1,18),(2,16),(3,14),(4,11)]:
    s=dst.styles[f'Heading {n}'];s.font.size=Pt(size);s.font.bold=True;s.paragraph_format.space_before=Pt(12)
    s.paragraph_format.keep_with_next=True
dst.styles['Caption'].font.size=Pt(9);dst.styles['Caption'].font.bold=False;dst.styles['Caption'].font.italic=True
dst.styles['Caption'].paragraph_format.space_after=Pt(9)
dst.styles['List Paragraph'].paragraph_format.left_indent=Cm(.4)
dst.styles['List Paragraph'].paragraph_format.first_line_indent=Cm(-.3)
lang=OxmlElement('w:lang');lang.set(qn('w:val'),'fr-FR');dst.styles['Normal'].element.get_or_add_rPr().append(lang)
sec.different_first_page_header_footer=True
h=sec.header.paragraphs[0];h.add_run('SEGULA TECHNOLOGIES  |  CAPTEUR EMBARQUÉ POUR VAC').font.size=Pt(8)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
f.add_run('VEHALSECU  |  ').font.size=Pt(8)
for name in ['PAGE','NUMPAGES']:
    if name=='NUMPAGES':f.add_run(' / ').font.size=Pt(8)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),name);f._p.append(field)

def para(text='',style=None):
    p=dst.add_paragraph(style=style)
    p.add_run(text);return p

def insert_img(p,blob,maxw=16,maxh=12):
    im=Image.open(BytesIO(blob));w,h=im.size
    wcm=min(maxw,maxh*w/h);p.add_run().add_picture(BytesIO(blob),width=Cm(wcm))
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
    p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(4)

def table(headers,rows,widths):
    t=dst.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for c,txt in zip(t.rows[0].cells,headers):c.text=txt
    for vals in rows:
        for c,txt in zip(t.add_row().cells,vals):c.text=txt
    for col,w in zip(t.columns,widths):col.width=Cm(w)
    pr=t._tbl.tblPr;borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    pr.append(borders)
    for ri,row in enumerate(t.rows):
        rp=row._tr.get_or_add_trPr();rp.append(OxmlElement('w:cantSplit'))
        if ri==0:rp.append(OxmlElement('w:tblHeader'))
        for c,w in zip(row.cells,widths):
            c.width=Cm(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');mar.append(e)
            cp.append(mar);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCEAF2' if ri==0 else ('F5F8FA' if ri%2==0 else 'FFFFFF'));cp.append(sh)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1.04
                for r in p.runs:r.font.size=Pt(9);r.bold=(ri==0)
    para().paragraph_format.space_after=Pt(2)
    return t

# Cover retains the SEGULA logo and original document identity.
p=para();insert_img(p,(work/'image2.png').read_bytes(),maxw=9.5,maxh=2)
p.paragraph_format.space_after=Pt(58)
p=para('Projets de Recherche et Innovation');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.runs[0].font.size=Pt(16)
p=para('Fiche de synthèse');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.runs[0].font.size=Pt(20);p.runs[0].bold=True
p.paragraph_format.space_after=Pt(30)
p=para('Capteur embarqué intelligent pour véhicules autonomes connectés','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=para('Implantation sur carte et extension vers une estimation résiliente','Subtitle');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after=Pt(38)
para('Projet VEHALSECU  |  CRAN et SEGULA Technologies').alignment=WD_ALIGN_PARAGRAPH.CENTER
para('Le démonstrateur embarqué réalise l’acquisition, la fusion de capteurs et l’export de l’état estimé. Cette fiche présente les travaux sur la carte, les performances rapportées et les extensions issues des travaux de thèse pour traiter les données corrompues.').paragraph_format.space_before=Pt(25)
dst.add_page_break()
para('Identification de l’opération','Heading 2')
table(['Rubrique','Information'],[
 ['Domaines de recherche','A1a, A2b, A3d'],
 ['Mots clés','Capteur embarqué ; STM32 ; EKF ; fusion GNSS et inertiel ; firmware temps réel ; ROS 2 ; confiance distribuée ; retour arrière ; Robust KalmanNet.'],
 ['Cadre scientifique','Thèse CIFRE VEHALSECU ; NGUYEN Quang Huy ; ANRT 2022/1732 ; CRAN et SEGULA Technologies.'],
 ['Volume horaire déclaré','Référence aux documents économiques de l’opération.']
],[4.8,11.2])
para('Sommaire','Title')
p=dst.add_paragraph();r=p.add_run();beg=OxmlElement('w:fldChar');beg.set(qn('w:fldCharType'),'begin');r._r.append(beg)
r=p.add_run();inst=OxmlElement('w:instrText');inst.set(qn('xml:space'),'preserve');inst.text=' TOC \\o "1-3" \\h \\z \\u ';r._r.append(inst)
r=p.add_run();sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');r._r.append(sep)
p.add_run('Sommaire actualisé à l’ouverture dans Word')
r=p.add_run();end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');r._r.append(end)
dst.add_page_break()

active=False;annex=False;counts=[0]*4
for el in src.element.body:
    if el.tag==qn('w:p'):
        p=Paragraph(el,src)
        if p.text=='Opération de R&D dans le cadre de l’activité de l’entreprise':active=True
        if not active:continue
        if not p.text and not p._p.xpath('.//a:blip'):continue
        if p.text.startswith('Annexes livrables'):
            annex=True
            ns=dst.add_section(WD_SECTION_START.NEW_PAGE);ns.orientation=WD_ORIENT.LANDSCAPE;ns.page_width=Cm(29.7);ns.page_height=Cm(21)
            ns.left_margin=Cm(2.5);ns.right_margin=Cm(2.25);ns.top_margin=Cm(2.5);ns.bottom_margin=Cm(2.5)
            ns.different_first_page_header_footer=False
        pics=p._p.xpath('.//a:blip')
        if pics:
            # Skip the source's decorative separator logo before the annex.
            if any(bl.get(qn('r:embed'))=='rId75' for bl in pics):continue
            np=para()
            for bl in pics:
                rid=bl.get(qn('r:embed'));part=src.part.related_parts.get(rid)
                if part:insert_img(np,part.blob,maxw=(15.8/len(pics)-.2 if len(pics)>1 else 15.8),maxh=9 if len(pics)==1 else 8)
            if p.text.strip():np.add_run(p.text)
            continue
        style=p.style.name if p.style.name in dst.styles else 'Normal'
        text=p.text
        if style.startswith('Heading'):
            lvl=int(style[-1]);counts[lvl-1]+=1
            for j in range(lvl,4):counts[j]=0
            text='.'.join(str(v) for v in counts[:lvl])+'  '+text
        np=para(text,style)
        if style=='Caption':np.alignment=WD_ALIGN_PARAGRAPH.CENTER
        if style=='List Paragraph':np.text='• '+text
        # Emphasize the provenance of the hypothetical resource figures.
        for lead in ['Architecture proposée.','Spécification de consolidation du firmware.','Exemple chiffré de dimensionnement uniquement.','Exemple mémoire calculé.','Mesures à produire sur carte.']:
            if text.startswith(lead):
                np.clear();np.add_run(lead).bold=True;np.add_run(text[len(lead):])
    elif el.tag==qn('w:tbl') and active:
        t=Table(el,src);rows=[[c.text for c in row.cells] for row in t.rows]
        n=len(rows[0]);available=24.95 if annex else 16.0
        if annex:
            # Retain the annex's substantive row and confidentiality context, removing template instructions.
            para('Annexe destinée à la relecture interne de la fiche de synthèse.').runs[0].italic=True
            data=rows[-1]
            for idx,val in enumerate(data):
                if idx<len(rows[0]):pass
            widths=[2.0,4.2,9.0,6.75,1.5,1.5] if n==6 else [available/n]*n
            if n==6:
                table(rows[0],[data],widths)
            else:table(rows[0],rows[-1:],widths)
        else:
            if n==4:widths=[4.1,2.1,3.2,6.6]
            elif n==2:widths=[6.1,9.9]
            elif n==3:
                if rows[0][0] in ['Indicateur','Fonction']:widths=[8.0,5.0,3.0]
                elif rows[0][0]=='Objectif de l’opération':widths=[5.5,2.5,8.0]
                else:widths=[5.2,4.7,6.1]
            else:widths=[available/n]*n
            table(rows[0],rows[1:],widths)

# Refresh the new TOC and page fields in Word.
uf=OxmlElement('w:updateFields');uf.set(qn('w:val'),'true');dst.settings.element.append(uf)
dst.core_properties.title='Capteur embarqué intelligent pour véhicules autonomes connectés'
dst.core_properties.subject='Implantation sur carte et extension vers une estimation résiliente'
dst.core_properties.author='NGUYEN Quang Huy'
dst.save(out)
print('Rebuilt',out,'paragraphs',len(dst.paragraphs),'tables',len(dst.tables))
